"""Bulk student import from an Excel (.xlsx) or CSV file."""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction

from .models import Classroom, Enrollment, Student

COLUMNS = (
    "admission_number",
    "first_name",
    "last_name",
    "class",
    "section",
    "guardian_name",
    "guardian_phone",
    "date_of_birth",
    "gender",
    "phone",
    "address",
    "area",
    "cnic",
    "email",
)
REQUIRED = ("admission_number", "first_name", "class", "section")
EXAMPLE_ROW = (
    "2026-001", "Ali", "Khan", "Grade 6", "A", "Imran Khan", "0300-1234567",
    "15/03/2014", "Male", "", "House 12, Street 4", "Gulshan", "", "",
)
# Common header spellings people use in their own sheets.
ALIASES = {
    "admission_no": "admission_number",
    "admission": "admission_number",
    "roll_number": "admission_number",
    "gr_number": "admission_number",
    "father_name": "guardian_name",
    "parent_name": "guardian_name",
    "father_phone": "guardian_phone",
    "parent_phone": "guardian_phone",
    "mobile": "guardian_phone",
    "dob": "date_of_birth",
    "grade": "class",
    "classroom": "class",
    "b_form": "cnic",
}
DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y")


@dataclass
class ImportResult:
    errors: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    created: int = 0
    updated: int = 0
    classes_created: int = 0


def template_csv():
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    writer.writerow(EXAMPLE_ROW)
    return buffer.getvalue()


def _header(value):
    key = str(value or "").strip().lower().replace(" ", "_").replace("-", "_").replace(".", "")
    return ALIASES.get(key, key)


def read_rows(uploaded):
    """Yield (row_number, dict) from an uploaded .xlsx or .csv file."""
    name = uploaded.name.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(uploaded, read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
    elif name.endswith(".csv"):
        text = uploaded.read().decode("utf-8-sig", errors="replace")
        rows = csv.reader(io.StringIO(text))
    else:
        raise ValueError("Upload an Excel (.xlsx) or CSV (.csv) file.")
    headers = None
    for number, values in enumerate(rows, start=1):
        if headers is None:
            headers = [_header(value) for value in values]
            continue
        if not any(str(value or "").strip() for value in values):
            continue
        yield number, dict(zip(headers, values))


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Date of birth '{value}' is not a date. Use DD/MM/YYYY.")


def _gender(value):
    text = _text(value).lower()
    if not text:
        return ""
    if text in ("m", "male", "boy"):
        return Student.Gender.MALE
    if text in ("f", "female", "girl"):
        return Student.Gender.FEMALE
    raise ValueError(f"Gender '{value}' should be Male or Female.")


def import_students(uploaded, year, create_classes=False):
    """Validate every row first; save only when the whole file is valid."""
    result = ImportResult()
    try:
        raw_rows = list(read_rows(uploaded))
    except ValueError as error:
        result.errors.append((None, str(error)))
        return result
    except Exception:
        result.errors.append((None, "The file could not be read. Save it again as .xlsx or .csv and retry."))
        return result
    if not raw_rows:
        result.errors.append((None, "The file has no student rows."))
        return result

    classrooms = {
        (c.name.lower(), c.section.lower()): c for c in Classroom.objects.filter(academic_year=year)
    }
    seen = set()
    for number, raw in raw_rows:
        data = {column: _text(raw.get(column)) for column in COLUMNS}
        missing = [column.replace("_", " ") for column in REQUIRED if not data[column]]
        if missing:
            result.errors.append((number, f"Missing {', '.join(missing)}."))
            continue
        if not data["last_name"] and " " in data["first_name"]:
            data["first_name"], data["last_name"] = data["first_name"].rsplit(" ", 1)
        if not data["last_name"]:
            result.errors.append((number, "Missing last name (or write the full name in first name)."))
            continue
        if data["admission_number"] in seen:
            result.errors.append((number, f"Admission number {data['admission_number']} appears twice in the file."))
            continue
        seen.add(data["admission_number"])
        try:
            data["date_of_birth"] = _date(raw.get("date_of_birth"))
            data["gender"] = _gender(raw.get("gender"))
        except ValueError as error:
            result.errors.append((number, str(error)))
            continue
        key = (data["class"].lower(), data["section"].lower())
        if key not in classrooms and not create_classes:
            result.errors.append(
                (number, f"Class {data['class']} {data['section']} does not exist in {year.name}.")
            )
            continue
        result.rows.append((number, data))

    if result.errors:
        return result

    with transaction.atomic():
        for number, data in result.rows:
            key = (data["class"].lower(), data["section"].lower())
            if key not in classrooms:
                classrooms[key] = Classroom.objects.create(
                    name=data["class"], section=data["section"], academic_year=year
                )
                result.classes_created += 1
            fields = {
                column: data[column]
                for column in COLUMNS
                if column not in ("admission_number", "class", "section")
            }
            student = Student.objects.filter(admission_number=data["admission_number"]).first()
            if student is None:
                student = Student(admission_number=data["admission_number"], is_active=True)
                result.created += 1
            else:
                result.updated += 1
            for name, value in fields.items():
                if value not in ("", None) or name in ("first_name", "last_name"):
                    setattr(student, name, value)
            try:
                student.full_clean()
            except ValidationError as error:
                transaction.set_rollback(True)
                messages = "; ".join(f"{k.replace('_', ' ')}: {' '.join(v)}" for k, v in error.message_dict.items())
                result.errors.append((number, messages))
                result.created = result.updated = result.classes_created = 0
                return result
            student.save()
            enrollment, _ = Enrollment.objects.get_or_create(
                student=student,
                academic_year=year,
                defaults={"classroom": classrooms[key]},
            )
            if enrollment.classroom_id != classrooms[key].pk or enrollment.status != Enrollment.Status.ACTIVE:
                enrollment.classroom = classrooms[key]
                enrollment.status = Enrollment.Status.ACTIVE
                enrollment.save()
    return result
